import re
from io import BytesIO
from typing import Any

from docx import Document


class BlueprintParser:
    """
    Parses company document blueprints.

    Supports TWO formats.

    ---------------------------------------------------------
    FORMAT 1 — RICH / RECOMMENDED
    ---------------------------------------------------------

    {
        "pan": {
            "document_type": "pan",
            "display_name": "PAN",
            "fields": [
                {
                    "name": "pan_name",
                    "original_name": "Name",
                    "type": "string",
                    "required": true,
                    "description": "Candidate name"
                }
            ]
        }
    }

    ---------------------------------------------------------
    FORMAT 2 — SIMPLE / LEGACY
    ---------------------------------------------------------

    {
        "pan": [
            "pan_name",
            "pan_dob",
            "pan_number"
        ]
    }

    The rich format is preserved.
    The parser NEVER treats:

        document_type
        display_name
        fields

    as extraction fields unless the user explicitly
    defines them inside the actual fields list.
    """

    # =========================================================
    # SANITIZE DOCUMENT / FIELD NAME
    # =========================================================

    @staticmethod
    def _sanitize_field_name(
        name: str,
    ) -> str:

        if name is None:
            return ""

        clean = re.sub(
            r"[^\w\s]",
            "",
            str(name),
        ).strip().lower()

        return re.sub(
            r"\s+",
            "_",
            clean,
        )

    # =========================================================
    # INFER TYPE FOR LEGACY BLUEPRINT
    # =========================================================

    @staticmethod
    def _infer_type(
        field_name: str,
    ) -> str:

        field_lower = (
            str(field_name)
            .lower()
        )

        # -----------------------------------------------------
        # DATE
        # -----------------------------------------------------

        if any(
            keyword in field_lower
            for keyword in [
                "dob",
                "date",
                "yop",
                "year",
                "start_date",
                "end_date",
            ]
        ):
            return "date"

        # -----------------------------------------------------
        # BOOLEAN
        # -----------------------------------------------------

        if any(
            keyword in field_lower
            for keyword in [
                "(y/n)",
                "accepted",
                "completed",
                "detected",
                "original",
                "acceptable",
            ]
        ):
            return "boolean"

        # -----------------------------------------------------
        # NUMBER
        # -----------------------------------------------------

        if any(
            keyword in field_lower
            for keyword in [
                "marks",
                "salary",
                "pay",
                "percentage",
                "amount",
                "age",
                "number",
            ]
        ):

            # Don't classify IDs / account numbers
            # as numeric values.
            if any(
                keyword in field_lower
                for keyword in [
                    "account_number",
                    "account_no",
                    "mobile_no",
                    "pan_number",
                    "aadhar_number",
                    "aadhaar_number",
                    "employee_id",
                    "uan_no",
                    "esi_no",
                ]
            ):
                return "string"

            return "number"

        # -----------------------------------------------------
        # ARRAY
        # -----------------------------------------------------

        if any(
            keyword in field_lower
            for keyword in [
                "skills",
                "subjects",
                "details",
                "particulars",
            ]
        ):
            return "array"

        return "string"

    # =========================================================
    # NORMALIZE RICH FIELD
    # =========================================================

    @classmethod
    def _normalize_rich_field(
        cls,
        field: Any,
    ) -> dict | None:
        """
        Normalize one field from the rich blueprint.

        Example input:

        {
            "name": "pan_name",
            "original_name": "Name",
            "type": "string",
            "required": true,
            "description": "Candidate name"
        }

        Output keeps the same semantic information.
        """

        # -----------------------------------------------------
        # FIELD AS STRING
        # -----------------------------------------------------

        if isinstance(
            field,
            str,
        ):

            field_name = (
                cls._sanitize_field_name(
                    field
                )
            )

            if not field_name:
                return None

            return {
                "name": field_name,
                "original_name": field,
                "type": cls._infer_type(
                    field
                ),
                "required": True,
                "description": (
                    f"Extracted value for {field}"
                ),
            }

        # -----------------------------------------------------
        # FIELD AS DICTIONARY
        # -----------------------------------------------------

        if isinstance(
            field,
            dict,
        ):

            raw_name = (
                field.get("name")
                or field.get("original_name")
            )

            if not raw_name:
                return None

            field_name = (
                cls._sanitize_field_name(
                    raw_name
                )
            )

            if not field_name:
                return None

            field_type = field.get(
                "type"
            )

            if not field_type:
                field_type = (
                    cls._infer_type(
                        str(raw_name)
                    )
                )

            required = field.get(
                "required",
                True,
            )

            description = field.get(
                "description"
            )

            if not description:
                description = (
                    f"Extracted value for "
                    f"{raw_name}"
                )

            return {
                "name": field_name,
                "original_name": str(
                    field.get(
                        "original_name",
                        raw_name,
                    )
                ),
                "type": str(
                    field_type
                ),
                "required": bool(
                    required
                ),
                "description": str(
                    description
                ),
            }

        return None

    # =========================================================
    # PARSE RAW JSON BLUEPRINT
    # =========================================================

    @classmethod
    def parse_raw_dict(
        cls,
        data: dict,
    ) -> dict:
        """
        Parse a company blueprint JSON.

        IMPORTANT:

        Rich format is handled correctly.

        We DO NOT iterate over:

            document_type
            display_name
            fields

        as if they were field names.

        Only the actual contents of:

            document_definition["fields"]

        become extraction parameters.
        """

        if not isinstance(
            data,
            dict,
        ):
            raise ValueError(
                "Blueprint must be a JSON object."
            )

        master_schemas = {}

        # =====================================================
        # DOCUMENT LOOP
        # =====================================================

        for (
            doc_name,
            document_definition,
        ) in data.items():

            # -------------------------------------------------
            # Normalize document key
            # -------------------------------------------------

            doc_key = (
                cls._sanitize_field_name(
                    doc_name
                )
            )

            if not doc_key:
                continue

            # =================================================
            # RICH FORMAT
            # =================================================

            if isinstance(
                document_definition,
                dict,
            ):

                document_type = (
                    document_definition.get(
                        "document_type"
                    )
                    or doc_key
                )

                display_name = (
                    document_definition.get(
                        "display_name"
                    )
                    or doc_name
                )

                raw_fields = (
                    document_definition.get(
                        "fields",
                        [],
                    )
                )

                if not isinstance(
                    raw_fields,
                    list,
                ):

                    raise ValueError(
                        f"Invalid fields definition "
                        f"for document '{doc_name}'. "
                        f"'fields' must be a list."
                    )

                field_definitions = []

                for field in raw_fields:

                    normalized_field = (
                        cls._normalize_rich_field(
                            field
                        )
                    )

                    if normalized_field:
                        field_definitions.append(
                            normalized_field
                        )

                master_schemas[
                    doc_key
                ] = {
                    "document_type": (
                        cls._sanitize_field_name(
                            str(
                                document_type
                            )
                        )
                    ),
                    "display_name": str(
                        display_name
                    ),
                    "fields": field_definitions,
                }

            # =================================================
            # LEGACY FORMAT
            # =================================================

            elif isinstance(
                document_definition,
                list,
            ):

                field_definitions = []

                for field in document_definition:

                    normalized_field = (
                        cls._normalize_rich_field(
                            field
                        )
                    )

                    if normalized_field:
                        field_definitions.append(
                            normalized_field
                        )

                master_schemas[
                    doc_key
                ] = {
                    "document_type": doc_key,
                    "display_name": str(
                        doc_name
                    ),
                    "fields": field_definitions,
                }

            # =================================================
            # INVALID
            # =================================================

            else:

                raise ValueError(
                    f"Invalid definition for "
                    f"document '{doc_name}'. "
                    f"Expected object or list."
                )

        # =====================================================
        # VALIDATION
        # =====================================================

        if not master_schemas:

            raise ValueError(
                "Blueprint does not contain "
                "any valid document definitions."
            )

        return master_schemas

    # =========================================================
    # PARSE DOCX
    # =========================================================

    @classmethod
    def parse_docx(
        cls,
        file_bytes: bytes,
    ) -> dict:
        """
        Parse a DOCX blueprint.

        Expected table structure:

        | Document Type | Fields |
        |---------------|--------|
        | PAN           | Name, DOB, PAN Number |
        | Aadhar        | Name, DOB, Aadhar Number |
        """

        if not file_bytes:
            raise ValueError(
                "DOCX file is empty."
            )

        doc = Document(
            BytesIO(
                file_bytes
            )
        )

        extracted_blueprint = {}

        # =====================================================
        # TABLES
        # =====================================================

        for table in doc.tables:

            for row in table.rows:

                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                if len(cells) < 2:
                    continue

                doc_type = cells[0]

                fields_text = cells[1]

                if not doc_type:
                    continue

                fields = [
                    field.strip()
                    for field in fields_text.split(
                        ","
                    )
                    if field.strip()
                ]

                if fields:

                    extracted_blueprint[
                        doc_type
                    ] = fields

        # =====================================================
        # TABLE FOUND
        # =====================================================

        if extracted_blueprint:

            return cls.parse_raw_dict(
                extracted_blueprint
            )

        # =====================================================
        # FALLBACK PARAGRAPHS
        # =====================================================

        paragraphs = [
            paragraph.text.strip()
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        ]

        if not paragraphs:

            raise ValueError(
                "No blueprint data found "
                "in DOCX file."
            )

        raise ValueError(
            "DOCX blueprint could not be parsed. "
            "Use a table with columns: "
            "Document Type and Fields."
        )


blueprint_parser = BlueprintParser()